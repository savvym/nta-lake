"""W3-5 DocxLoader behavioral tests."""

from __future__ import annotations

import base64
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest

# Minimal 1x1 PNG (transparent) — 用 base64 嵌入
_PNG_1x1_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNgYAAAAAMAASsJTYQAAAAASUVORK5CYII="
)
_PNG_BYTES = base64.b64decode(_PNG_1x1_B64)


class _StubBlobStore:
    """In-memory BlobStore stub: 支持 async get + async put。"""

    def __init__(self, data: dict[str, bytes]) -> None:
        self._data = dict(data)
        self.put_count = 0

    async def get(self, sha: str) -> bytes:
        return self._data[sha]

    async def put(self, stream: BytesIO, declared_size: int) -> Any:
        self.put_count += 1
        # 返一个简单对象，模拟 BlobPutResult
        return SimpleNamespace(
            sha256="b" * 64,
            size=declared_size,
            storage_key="stub-key",
            deduplicated=False,
        )


def _make_ctx(blob_store: Any | None) -> Any:
    return SimpleNamespace(
        blob_store=blob_store,
        logger=None,
        metrics=None,
        secrets=None,
        cancel_event=None,
        llm=None,
    )


def _build_sample_docx() -> bytes:
    """运行时构造一个含 paragraph + image 的 .docx bytes。"""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello docx W3-5 test")
    doc.add_paragraph("Second paragraph")
    doc.add_picture(BytesIO(_PNG_BYTES))
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def test_docx_auto_registered() -> None:
    from dataplat_core.loaders import DocxLoader, LoaderRegistry

    assert "docx" in LoaderRegistry.list_names()
    assert LoaderRegistry.get("docx") is DocxLoader


def test_docx_load_happy() -> None:
    from dataplat_core.loaders import DocxLoader

    sha = "a" * 64
    docx_bytes = _build_sample_docx()
    ctx = _make_ctx(_StubBlobStore({sha: docx_bytes}))

    loader = DocxLoader()
    result = loader.load(sha, {}, ctx)

    assert result.total_count == 1
    row = result.rows[0]
    assert "Hello docx" in row.text
    assert row.stats["format"] == "docx"
    assert row.stats["paragraph_count"] >= 1
    assert row.stats["image_count"] >= 1
    assert len(row.images) >= 1
    assert row.images[0]["blob_sha"] == "b" * 64
    assert row.images[0]["content_type"].startswith("image/")


def test_docx_requires_blob_store() -> None:
    from dataplat_core.loaders import DocxLoader

    loader = DocxLoader()
    with pytest.raises(ValueError, match="ctx.blob_store"):
        loader.load("a" * 64, {}, _make_ctx(None))
