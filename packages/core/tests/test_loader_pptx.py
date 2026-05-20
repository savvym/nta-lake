"""W3-5 PptxLoader behavioral tests."""

from __future__ import annotations

import base64
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest

# Minimal 1x1 PNG (transparent) — 用 base64 嵌入
_PNG_1x1_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNgYAAAAAMAASsJTYQAAAAASUVORK5CYII="
)
_PNG_BYTES = base64.b64decode(_PNG_1x1_B64)


class _StubBlobStore:
    """In-memory BlobStore stub: 支持 async get + async put。"""

    def __init__(self, data: dict[str, bytes]) -> None:
        self._data = dict(data)
        self.put_count = 0

    async def get(self, sha: str) -> bytes:
        return self._data[sha]

    async def put(self, stream: BytesIO, declared_size: int) -> Any:
        self.put_count += 1
        return SimpleNamespace(
            sha256="b" * 64,
            size=declared_size,
            storage_key="stub-key",
            deduplicated=False,
        )


def _make_ctx(blob_store: Any | None) -> Any:
    return SimpleNamespace(
        blob_store=blob_store,
        logger=None,
        metrics=None,
        secrets=None,
        cancel_event=None,
        llm=None,
    )


def _build_sample_pptx() -> bytes:
    """运行时构造一个含 1 张 slide + text + image 的 .pptx bytes。"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide_layout = prs.slide_layouts[5]  # Title Only
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Hello pptx W3-5"
    slide.shapes.add_picture(BytesIO(_PNG_BYTES), Inches(1), Inches(1), width=Inches(1), height=Inches(1))
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def test_pptx_auto_registered() -> None:
    from dataplat_core.loaders import LoaderRegistry, PptxLoader

    assert "pptx" in LoaderRegistry.list_names()
    assert LoaderRegistry.get("pptx") is PptxLoader


def test_pptx_load_happy() -> None:
    from dataplat_core.loaders import PptxLoader

    sha = "a" * 64
    pptx_bytes = _build_sample_pptx()
    ctx = _make_ctx(_StubBlobStore({sha: pptx_bytes}))

    loader = PptxLoader()
    result = loader.load(sha, {}, ctx)

    assert result.total_count == 1
    row = result.rows[0]
    assert "Hello pptx" in row.text
    assert row.stats["format"] == "pptx"
    assert row.stats["slide_count"] == 1
    assert row.stats["image_count"] >= 1
    assert len(row.images) >= 1
    assert row.images[0]["blob_sha"] == "b" * 64
    assert row.images[0]["content_type"].startswith("image/")


def test_pptx_requires_blob_store() -> None:
    from dataplat_core.loaders import PptxLoader

    loader = PptxLoader()
    with pytest.raises(ValueError, match="ctx.blob_store"):
        loader.load("a" * 64, {}, _make_ctx(None))
