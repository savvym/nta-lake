"""PptxLoader：Bronze .pptx blob → SilverRow（via python-pptx）。

stats:
  - format: "pptx"
  - slide_count: int
  - image_count: int
  - char_count: int

图片写 blob_store；images 列含 {filename, blob_sha, content_type}。
"""

from __future__ import annotations

import asyncio
from io import BytesIO
from typing import Any

from dataplat_core.domain.types import SHA256
from dataplat_core.protocols.loader import LoadResult, SilverRow
from dataplat_core.protocols.runcontext import RunContext

_MSO_PICTURE = 13  # MSO_SHAPE_TYPE.PICTURE


class PptxLoader:
    name: str = "pptx"
    version: str = "0.1"
    input_subtype: str = "pptx"
    output_schema_id: str = "silver-text-v1"

    def load(
        self,
        bronze_blob_sha: SHA256,
        config: dict[str, Any],
        ctx: RunContext,
    ) -> LoadResult:
        del config  # 暂不接受 config 选项
        blob_store = getattr(ctx, "blob_store", None)
        if blob_store is None:
            raise ValueError(
                "loader pptx 需 ctx.blob_store；当前 RunContext 未注入"
            )

        async def _run() -> LoadResult:
            from pptx import Presentation

            data = await blob_store.get(bronze_blob_sha)
            if not isinstance(data, bytes):
                chunks = []
                async for chunk in data:
                    chunks.append(chunk)
                data = b"".join(chunks)

            prs = Presentation(BytesIO(data))

            slide_texts: list[str] = []
            images: list[dict[str, Any]] = []
            pic_idx = 0

            for slide_idx, slide in enumerate(prs.slides):
                parts: list[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        shape_text = shape.text_frame.text.strip()
                        if shape_text:
                            parts.append(shape_text)
                    if shape.shape_type == _MSO_PICTURE:
                        img = shape.image
                        img_bytes = img.blob
                        ct = img.content_type or "image/unknown"
                        try:
                            filename = img.filename
                        except Exception:
                            filename = None
                        if not filename:
                            filename = f"slide-{slide_idx}-pic-{pic_idx}.{img.ext}"
                        put_res = await blob_store.put(
                            BytesIO(img_bytes), declared_size=len(img_bytes)
                        )
                        images.append({
                            "filename": filename,
                            "blob_sha": put_res.sha256,
                            "content_type": ct,
                        })
                        pic_idx += 1

                slide_text = "\n".join(parts)
                if slide_text:
                    slide_texts.append(f"\n\n[slide {slide_idx + 1}]\n\n{slide_text}")
                else:
                    slide_texts.append(f"\n\n[slide {slide_idx + 1}]\n\n")

            text = "".join(slide_texts).strip()

            row = SilverRow(
                text=text,
                images=images,
                source_ref={
                    "blob_sha": bronze_blob_sha,
                    "loader": "pptx",
                    "loader_version": "0.1",
                },
                stats={
                    "format": "pptx",
                    "slide_count": len(prs.slides),
                    "image_count": len(images),
                    "char_count": len(text),
                },
                lineage_ops=[],
            )

            return LoadResult(
                rows=[row],
                total_count=1,
                notes=f"pptx slides={len(prs.slides)}, images={len(images)}",
            )

        return asyncio.run(_run())
